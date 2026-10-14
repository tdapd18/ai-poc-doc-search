"""가상 PDF / Excel / PowerPoint 샘플 파일 생성기 (개발 전용).

앱 실행에는 필요 없다(앱의 추출기는 표준 라이브러리만 사용). 이 스크립트는 샘플 파일을 다시 만들 때만 쓴다.
  pip install -r requirements-dev.txt
  python tools/make_sample_files.py
모든 내용은 가상이며, 문서 메타데이터(ID/권한 등)는 파일 옆의 `<파일명>.meta.json` 에 기록한다(가상 ACL 대용).
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent / "data" / "shared_folder"


def meta(path: Path, **kw) -> None:
    base = dict(allowed_users=[], version="1.0", department="계리평가팀")
    base.update(kw)
    Path(str(path) + ".meta.json").write_text(json.dumps(base, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


# ---------------------------------------------------------------- PDF
def make_pdf() -> None:
    p = ROOT / "계리평가" / "결산검증" / "결산_내부통제_점검보고서.pdf"
    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    pages = [
        ["IFRS17 결산 내부통제 점검 보고서 (2026년 9월)", "",
         "1. 점검 목적", "월말 결산 산출물의 입력, 모델 실행, 검증 절차가 기준에 따라 수행되었는지 확인한다.",
         "2. 점검 범위", "RAFM 모델 실행 이력, 가정 버전 통제, GoC 롤포워드 검증, 재무팀 인터페이스 대사",
         "3. 점검 기간", "2026-10-02 ~ 2026-10-07"],
        ["4. 점검 결과", "",
         "입력 데이터 대사: 적합 (계약 건수 전월 대비 +0.8%)",
         "모델 실행 로그 점검: RAFM 실행 오류 2건(E101, E207) 발생, 재실행으로 모두 조치 완료",
         "가정 버전 통제: 적합 (가정 v3.0 일괄 적용 확인)",
         "롤포워드 검증: 조건부 적합 (GoC-C07 CSM 변동률 6.2% 초과, 변동 사유서 작성)",
         "재무팀 인터페이스 대사: 적합 (차이 0원)"],
        ["5. 개선 과제", "",
         "가. RAFM 모델 실행 오류 재발 방지: 입력 파일 도착 여부를 실행 전에 자동 확인한다.",
         "나. 메모리 부족 에러 예방: 상품군별 분할 실행을 표준 절차로 지정한다.",
         "다. 변동 사유서 서식을 표준화하고 결산 검증 결과 보고서에 첨부한다.", "",
         "6. 결재", "작성: 박계리(가상)  검토: 팀장(가상)  승인: 부문장(가상)"],
    ]
    c = canvas.Canvas(str(p), pagesize=A4, pageCompression=0)
    for lines in pages:
        c.setFont("HYSMyeongJo-Medium", 12)
        y = 780
        for line in lines:
            if line:
                c.drawString(50, y, line)
            y -= 22
        c.showPage()
    c.save()
    meta(p, doc_id="SF-101", title="결산 내부통제 점검 보고서 (2026년 9월)", category="결산", version="1.0",
         created_at="2026-10-07", modified_at="2026-10-08", allowed_groups=["actuarial", "finance"])


# ---------------------------------------------------------------- Excel
def make_xlsx_model() -> None:
    p = ROOT / "계리평가" / "모델검증" / "모델별_검증결과_집계_2026-09.xlsx"
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.append(["모델", "실행횟수", "오류건수", "재실행건수", "검증판정"])
    ws.append(["장기보장 모델", 12, 0, 0, "적합"])
    ws.append(["연금저축 모델", 8, 1, 1, "적합 (E101 재실행 후)"])
    ws.append(["일반보장 모델", 15, 1, 1, "적합 (E207 분할 실행 후)"])
    g = wb.create_sheet("GoC별변동")
    g.append(["GoC", "기초 CSM", "기말 CSM", "변동률", "사유"])
    g.append(["GoC-A01", 1200, 1197, "-0.3%", "정상 범위"])
    g.append(["GoC-B02", 860, 873, "1.5%", "정상 범위"])
    g.append(["GoC-C07", 640, 680, "6.2%", "신계약 집중 인수 - 변동 사유서 작성 대상"])
    e = wb.create_sheet("오류내역")
    e.append(["실행ID", "오류코드", "조치내용"])
    e.append(["RUN-0930-02", "E101", "입력 파일 미도착으로 모델 실행 실패, 재전송 후 재실행"])
    e.append(["RUN-0930-04", "E207", "메모리 부족 에러, 상품군별 분할 실행으로 해결"])
    wb.save(p)
    meta(p, doc_id="SF-102", title="모델별 검증결과 집계 (2026-09)", category="RAFM", version="1.2",
         created_at="2026-10-05", modified_at="2026-10-06", allowed_groups=["actuarial"])


def make_xlsx_finance() -> None:
    p = ROOT / "재무공유" / "결산일정" / "재무_결산_마감_체크시트.xlsx"
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "마감체크"
    ws.append(["단계", "담당", "기한", "상태"])
    ws.append(["결산용 계약 데이터 마감", "재무팀", "10월 30일", "진행 중"])
    ws.append(["계리 산출 결과 수령", "재무팀", "11월 6일", "대기"])
    ws.append(["인터페이스 금액 대사", "재무팀", "11월 8일", "대기"])
    ws.append(["재무제표 확정", "재무팀", "11월 10일", "대기"])
    d = wb.create_sheet("인터페이스대사")
    d.append(["계정", "계리 산출", "회계 반영", "차이"])
    d.append(["보험계약부채", 98200, 98200, 0])
    d.append(["CSM 상각", 1530, 1530, 0])
    wb.save(p)
    meta(p, doc_id="SF-201", title="재무 결산 마감 체크시트", category="결산", version="1.0", department="재무팀",
         created_at="2026-10-01", modified_at="2026-10-12", allowed_groups=["finance", "actuarial"])


# ---------------------------------------------------------------- PowerPoint
def make_pptx() -> None:
    p = ROOT / "계리평가" / "교육자료" / "월말결산_킥오프_자료.pptx"
    p.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "2026년 10월 월말 결산 킥오프"
    s.placeholders[1].text = "계리평가팀 · 재무팀 합동 (가상 자료)"

    def bullets(title, items, notes=None):
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = title
        tf = sl.placeholders[1].text_frame
        tf.text = items[0]
        for it in items[1:]:
            tf.add_paragraph().text = it
        if notes:
            sl.notes_slide.notes_text_frame.text = notes

    bullets("결산 일정", ["10월 30일 결산용 계약 데이터 마감", "11월 2일 RAFM 모델 실행 시작",
                        "11월 4일 모델 실행 결과 확인 및 결산 검증", "11월 6일 결산 검증 결과 보고서 제출"])
    bullets("RAFM 모델 실행 오류 사전 점검", ["입력 파일 날짜 접미사와 도착 여부 확인(E101 예방)",
                                          "가정 테이블 버전 확인", "동시 실행 작업 수 확인(E207 메모리 부족 예방)"],
            notes="지난 결산에서 E101, E207 오류가 각 1회 발생했다. 실행 전 점검 항목을 체크리스트로 확인한다.")
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    sl.shapes.title.text = "GoC 변동 분석 기준"
    rows = [("항목", "기준"), ("CSM 변동률", "전월 대비 ±5% 초과 시 변동 사유서"),
            ("변동 금액", "5억원 이상 우선 분석"), ("롤포워드 오차", "1백만원 이내")]
    tbl = sl.shapes.add_table(len(rows), 2, Inches(0.7), Inches(1.8), Inches(8.6), Inches(2.4)).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            tbl.cell(i, j).text = v
    bullets("역할 분담", ["계리평가팀: 모델 실행, 결산 검증, 변동 분석", "재무팀: 데이터 마감, 인터페이스 대사, 재무제표 확정"])
    prs.save(p)
    meta(p, doc_id="SF-103", title="월말 결산 킥오프 자료", category="결산", version="1.0",
         created_at="2026-09-28", modified_at="2026-10-02", allowed_groups=["actuarial", "finance"])


def make_txt() -> None:
    p = ROOT / "재무공유" / "인터페이스" / "계리_재무_인터페이스_대사_메모.txt"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text("""---
doc_id: SF-202
title: 계리 재무 인터페이스 대사 메모
category: 결산
version: 1.1
created_at: 2026-08-30
modified_at: 2026-09-20
department: 재무팀
allowed_users: []
allowed_groups: [finance]
---
# 인터페이스 대사 메모 (재무팀 내부)

가상의 교육용 메모입니다. 계리 산출 금액과 회계 시스템 반영 금액의 대사 방법을 정리합니다.

## 대사 원칙

인터페이스 금액 차이는 0원이어야 합니다. 차이가 발생하면 반올림 규칙 차이인지, 입력 시점 차이인지 먼저 확인하고 계리평가팀에 원인 확인을 요청합니다.

## 요청 절차

차이 금액, 계정, 기준일을 메일로 정리하여 계리평가팀 담당자에게 전달하고 회신 후 재대사합니다.
""", encoding="utf-8")


if __name__ == "__main__":
    make_pdf()
    make_xlsx_model()
    make_xlsx_finance()
    make_pptx()
    make_txt()
    print("sample files written under", ROOT)
