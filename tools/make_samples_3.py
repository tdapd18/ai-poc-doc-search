"""Generate FICTIONAL multi-format sample documents into data/3안_다형식샘플/.

Dev-only (needs reportlab + a Korean TTF such as malgun.ttf for the PDF). Not used by the app.
Run from the kit root: python tools/make_samples_3.py
"""
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent.parent / 'data' / '3안_다형식샘플'
OUT.mkdir(parents=True, exist_ok=True)


def make_docx():
    d = Document()
    d.core_properties.title = '교육용 가상 결산 승인 절차'
    d.add_heading('교육용 가상 결산 승인 절차', level=1)
    d.add_paragraph('- 버전: 2026-10 교육용')
    d.add_paragraph('- 분류: 가상 결산 절차')
    d.add_paragraph('- 담당자: 가상 결산팀')
    d.add_heading('결산 승인 절차', level=2)
    d.add_paragraph('가상 결산 결과는 담당자 검토, 팀장 확인, 승인자 서명 순서로 승인한다. 승인 전에는 결과를 확정하지 않는다.')
    d.add_heading('승인 일정표', level=2)
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([('단계', '기한', '담당'), ('담당자 검토', '마감 D+1', '가상 결산팀'), ('팀장 확인', '마감 D+2', '가상 팀장')]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.save(OUT / '3안_결산승인절차_가상.docx')


def make_xlsx():
    wb = Workbook()
    wb.properties.title = '교육용 가상 상품 한도표'
    ws = wb.active
    ws.title = '한도표'
    ws.append(['상품코드', '상품명', '한도구분', '한도값', '담당'])
    ws.append(['A100', '가상 종신 상품', '월 증감 한도', '5.0', '가상 상품팀'])
    ws.append(['B200', '가상 연금 상품', '월 증감 한도', '8.0', '가상 상품팀'])
    ws.append(['C300', '가상 건강 상품', '월 증감 한도', '3.5', '가상 상품팀'])
    wb.save(OUT / '3안_상품한도표_가상.xlsx')


def make_pdf():
    pdfmetrics.registerFont(TTFont('Malgun', r'C:\Windows\Fonts\malgun.ttf'))
    c = canvas.Canvas(str(OUT / '3안_데이터백업안내_가상.pdf'), pagesize=A4)
    c.setTitle('교육용 가상 데이터 백업 안내')
    pages = [['데이터 백업 주기', '가상 결산 데이터는 매일 1회 백업하며 월 마감일에는 추가 백업을 수행한다.'],
             ['백업 보관 기간', '가상 백업 파일은 90일 보관 후 폐기한다. 보관 중 접근은 가상 보안팀 승인이 필요하다.']]
    for heading, body in pages:
        c.setFont('Malgun', 16)
        c.drawString(60, 780, heading)
        c.setFont('Malgun', 11)
        c.drawString(60, 750, body)
        c.showPage()
    c.save()


def make_hwpx():
    ns = 'http://www.hancom.co.kr/hwpml/2011/paragraph'

    def p(text):
        return f'<hp:p><hp:run><hp:t>{text}</hp:t></hp:run></hp:p>'

    lines = ['RAFM 실행 환경 점검 메모', '버전: 2026-10 교육용', '분류: 가상 실행 환경', '담당자: 가상 운영팀',
             'RAFM 실행 전 가상 환경 점검 항목은 디스크 여유 공간, 입력 폴더 접근 권한, 기준일 설정 확인이다.',
             '점검 결과는 가상 운영 로그에 기록하고 이상이 있으면 가상 운영팀에 전달한다.']
    xml = f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="{ns}">' + ''.join(map(p, lines)) + '</hs:sec>'
    with zipfile.ZipFile(OUT / '3안_RAFM환경점검메모_가상.hwpx', 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('mimetype', 'application/hwp+zip')
        z.writestr('Contents/section0.xml', xml)


if __name__ == '__main__':
    make_docx(); make_xlsx(); make_pdf(); make_hwpx()
    print('created in', OUT)
