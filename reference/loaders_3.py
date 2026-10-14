"""PoC 3 multi-format loaders: md / docx / xlsx / pdf / hwpx -> core_3 chunk dicts.

All parsing is in-memory (no files written). Chunk keys match core_3.chunks_from_text,
plus 'format'. Missing metadata in non-Markdown files is shown as '미상'.
"""
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from reference.core_3 import MAX_DOC_BYTES, chunks_from_text

MAX_BINARY_BYTES = 5 * 1024 * 1024
MAX_UNZIPPED_BYTES = 30 * 1024 * 1024
MAX_XLSX_ROWS = 5000
MAX_PDF_PAGES = 300
SUPPORTED = ('md', 'docx', 'xlsx', 'pdf', 'hwpx')
UNKNOWN = '미상'
META_KEYS = {'버전': 'version', '분류': 'category', '담당자': 'owner'}


def _meta(text: str):
    """Return (key, value) when a line looks like '- 버전: x' / '버전: x'."""
    m = re.match(r'^[-\s]*(버전|분류|담당자)\s*[:：]\s*(.+)$', text.strip())
    return (META_KEYS[m.group(1)], m.group(2).strip()) if m else None


def _build(name: str, fmt: str, blocks: list[tuple], default_title: str = '') -> list[dict]:
    """blocks: ('title'|'section'|'text', str) or ('meta', (key, value))."""
    title = ''
    meta = {'version': UNKNOWN, 'category': UNKNOWN, 'owner': UNKNOWN}
    section, content, chunks = '', [], []

    def flush():
        body = ' '.join(content).strip()
        if body:
            chunks.append({'title': title or default_title or Path(name).stem,
                           **meta, 'file': name, 'section': section or '본문',
                           'text': body, 'source': f'{name}#{section or "본문"}',
                           'format': fmt})

    for kind, value in blocks:
        if kind == 'meta':
            meta[value[0]] = value[1]
        elif kind == 'title' and not title:
            title = value
        elif kind in ('title', 'section'):
            flush()
            section, content = value, []
        elif kind == 'text' and value.strip():
            content.append(value.strip())
    flush()
    if not chunks:
        raise ValueError(f'{name}: 읽을 수 있는 본문 텍스트가 없습니다.')
    # metadata may appear after the first section; apply the final values everywhere
    for c in chunks:
        c.update(meta)
        c['title'] = title or default_title or Path(name).stem
    return chunks


def _load_docx(name: str, data: bytes) -> list[dict]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f'{name}: 유효한 .docx 파일이 아닙니다.') from exc
    blocks = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else '').lower()
            if style == 'title' or style == 'heading 1':
                blocks.append(('title', text))
            elif style.startswith('heading'):
                blocks.append(('section', text))
            elif _meta(text):
                blocks.append(('meta', _meta(text)))
            else:
                blocks.append(('text', text))
        elif tag == 'tbl':
            for row in Table(child, doc).rows:
                cells = []
                for cell in row.cells:
                    t = cell.text.strip()
                    if t and (not cells or cells[-1] != t):
                        cells.append(t)
                if cells:
                    blocks.append(('text', ' | '.join(cells)))
    return _build(name, 'docx', blocks, (doc.core_properties.title or '').strip())


def _load_xlsx(name: str, data: bytes) -> list[dict]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f'{name}: 유효한 .xlsx 파일이 아닙니다.') from exc
    title = (wb.properties.title or '').strip() or Path(name).stem
    chunks = []
    for ws in wb.worksheets:
        header = None
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            if r > MAX_XLSX_ROWS:
                break
            cells = ['' if v is None else str(v).strip() for v in row]
            if not any(cells):
                continue
            if header is None:
                header = cells
                continue
            parts = [f'{h or f"열{i + 1}"}: {v}' for i, (h, v) in enumerate(zip(header, cells)) if v]
            if parts:
                section = f'{ws.title} {r}행'
                chunks.append({'title': title, 'version': UNKNOWN, 'category': UNKNOWN,
                               'owner': UNKNOWN, 'file': name, 'section': section,
                               'text': ', '.join(parts), 'source': f'{name}#{section}',
                               'format': 'xlsx'})
    wb.close()
    if not chunks:
        raise ValueError(f'{name}: 헤더 아래에 읽을 수 있는 데이터 행이 없습니다.')
    return chunks


def _load_pdf(name: str, data: bytes) -> list[dict]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError(f'{name}: 암호가 걸린 PDF는 지원하지 않습니다.')
        pages = reader.pages[:MAX_PDF_PAGES]
        texts = [(p.extract_text() or '').strip() for p in pages]
        info_title = ((reader.metadata or {}).get('/Title') or '').strip()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f'{name}: 유효한 PDF 파일이 아닙니다.') from exc
    blocks = []
    for n, text in enumerate(texts, 1):
        if text:
            blocks.append(('section', f'{n}쪽'))
            blocks.append(('text', re.sub(r'\s+', ' ', text)))
    if not blocks:
        raise ValueError(f'{name}: 추출 가능한 텍스트가 없습니다. 스캔본 PDF는 OCR 미지원입니다.')
    return _build(name, 'pdf', blocks, info_title)


def _load_hwpx(name: str, data: bytes) -> list[dict]:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ValueError(f'{name}: 유효한 .hwpx 파일이 아닙니다.') from exc
    with zf:
        if sum(i.file_size for i in zf.infolist()) > MAX_UNZIPPED_BYTES:
            raise ValueError(f'{name}: 압축 해제 크기가 제한을 초과했습니다.')
        sections = sorted((n for n in zf.namelist()
                           if re.fullmatch(r'Contents/section\d+\.xml', n)),
                          key=lambda n: int(re.search(r'(\d+)\.xml', n).group(1)))
        if not sections:
            raise ValueError(f'{name}: HWPX 본문(Contents/section*.xml)을 찾지 못했습니다.')
        blocks = []
        for idx, member in enumerate(sections, 1):
            try:
                root = ET.fromstring(zf.read(member))
            except ET.ParseError as exc:
                raise ValueError(f'{name}: HWPX 본문 XML이 손상되었습니다.') from exc
            paras = []
            for p in root.iter():
                if p.tag.rsplit('}', 1)[-1] == 'p':
                    text = ''.join(t.text or '' for t in p.iter() if t.tag.rsplit('}', 1)[-1] == 't').strip()
                    if text:
                        paras.append(text)
            for start in range(0, len(paras), 8):
                group = paras[start:start + 8]
                blocks.append(('section', f'섹션{idx} 문단 {start + 1}-{start + len(group)}'))
                for text in group:
                    blocks.append(('meta', _meta(text)) if _meta(text) else ('text', text))
    return _build(name, 'hwpx', blocks)


def load_bytes(name: str, data: bytes) -> list[dict]:
    """Parse one uploaded file in memory. Raises ValueError with a user-facing message."""
    ext = Path(name).suffix.lower().lstrip('.')
    if ext == 'hwp':
        raise ValueError(f'{name}: 구형 .hwp(바이너리)는 지원하지 않습니다. 한글에서 .hwpx로 저장해 주세요.')
    if ext not in SUPPORTED:
        raise ValueError(f'{name}: 지원하지 않는 형식입니다. ({", ".join("." + e for e in SUPPORTED)})')
    if ext == 'md':
        if len(data) > MAX_DOC_BYTES:
            raise ValueError(f'{name}: 파일이 1MB를 초과했습니다.')
        try:
            chunks = chunks_from_text(name, data.decode('utf-8-sig'))
        except UnicodeError as exc:
            raise ValueError(f'{name}: UTF-8 인코딩 문서가 아닙니다.') from exc
        return [{**c, 'format': 'md'} for c in chunks]
    if len(data) > MAX_BINARY_BYTES:
        raise ValueError(f'{name}: 파일이 5MB를 초과했습니다.')
    return {'docx': _load_docx, 'xlsx': _load_xlsx, 'pdf': _load_pdf, 'hwpx': _load_hwpx}[ext](name, data)


def load_folder(folder: Path) -> tuple[list[dict], list[dict]]:
    """Load every supported file in a folder. Returns (chunks, per-file status rows)."""
    chunks, status = [], []
    for path in sorted(p for p in folder.iterdir() if p.suffix.lower().lstrip('.') in SUPPORTED + ('hwp',)):
        try:
            got = load_bytes(path.name, path.read_bytes())
            chunks.extend(got)
            status.append({'파일': path.name, '형식': got[0]['format'], '상태': '성공', '문단 수': len(got), '메시지': ''})
        except ValueError as exc:
            status.append({'파일': path.name, '형식': path.suffix.lstrip('.').lower(), '상태': '실패', '문단 수': 0, '메시지': str(exc)})
    return chunks, status
